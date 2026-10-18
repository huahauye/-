#!/usr/bin/env python3
"""OA docx mechanical pre-submission QA. This checks obvious conflicts only, not legal correctness.
Usage: python qa_docx.py --docx FILE --metadata assets/case_metadata.example.json
Requires: python-docx.
"""
from __future__ import annotations
import argparse, json, re, sys
from pathlib import Path
from docx import Document


def normalize(s):
    return re.sub(r"[\s.·‐‑-]", "", str(s or "")).upper()


def iter_all_doc_text(document):
    for p in document.paragraphs:
        yield p.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text
    for section in document.sections:
        for block in (section.header, section.footer):
            for p in block.paragraphs:
                yield p.text
            for table in block.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield cell.text


def check(docx_path, meta_path):
    d = Document(docx_path)
    meta = json.loads(Path(meta_path).read_text(encoding="utf-8"))
    lines = list(iter_all_doc_text(d))
    full = "\n".join(lines)
    errors, warnings = [], []
    target = normalize(meta.get("application_number"))
    if not target:
        errors.append("元数据application_number缺失，不能进行案件号核验")
    # Look only within the table's application-number field to avoid flagging cited D1 document publication numbers.
    rows = []
    for t in d.tables:
        for r in t.rows:
            text = " ".join(c.text for c in r.cells)
            if "申请号或专利号" in text:
                rows.append(text)
    if rows:
        seen = re.findall(r"(?<!\d)\d{12}\s*\.?\s*[0-9Xx](?!\d)", " ".join(rows))
        unique = sorted(set(normalize(s) for s in seen))
        if not unique:
            warnings.append("申请号栏无法用常规模式解析，请人工对照")
        for value in unique:
            if value != target:
                errors.append(f"申请号栏存在与案件元数据不一致的编号：{value}；预期{target}")
    else:
        warnings.append("未发现国知局表格申请号栏（若仅为附页可忽略）")
    if meta.get("title") and meta["title"].strip() not in full:
        warnings.append("整份文件没有发现元数据中的发明名称；确认是否为仅正文附页")
    if meta.get("applicant") and meta["applicant"].strip() not in full:
        warnings.append("整份文件没有发现申请人全称；确认是否为仅正文附页")
    num=meta.get("claims_total")
    if isinstance(num,int) and num>0:
        # Heuristic check for claims ranges and numbering introduced in creativity sections; exclude previous-claim discussion.
        for s in lines:
            ss=re.sub(r"\s+", "", s)
            if "修改后的权利要求" in ss and "具有创造性" in ss:
                mm=re.search(r"修改后的权利要求(\d+)[—－~～-](\d+)具有创造性",ss)
                if mm:
                    mx=int(mm.group(2))
                    if mx>num:
                        errors.append(f"创造性段落的权项范围写至{mx}，超过元数据总数{num}：{ss[:110]}")
        # Last-page claims list may continue with stale numbering; manual reminder.
        warnings.append(f"需人工核对权利要求1—{num}正文、替换页、修改对照页及引用关系，自动程序无法证明一致")
    if re.search(r"[，,]{2,}|[。．]{2,}",full):
        warnings.append("发现连续重复标点，需人工核查")
    if "用于用于" in full:
        warnings.append("发现可疑重复词“用于用于”")
    if "原权利要求" in full and "修改对照" not in full:
        warnings.append("文档有原权项修改，但未检测到修改对照信息")
    return {"file":str(docx_path),"error_count":len(errors),"errors":list(dict.fromkeys(errors)),"warnings":list(dict.fromkeys(warnings)),"note":"机械核查不等于法律合规、修改超范围或新颖性创造性保证"}


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--docx",required=True)
    ap.add_argument("--metadata",required=True)
    ap.add_argument("--out-json")
    args=ap.parse_args()
    result=check(args.docx,args.metadata)
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if args.out_json:Path(args.out_json).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return 1 if result["error_count"] else 0

if __name__=="__main__": sys.exit(main())
