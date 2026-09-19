from __future__ import annotations

import logging
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import fitz
import magic
from docx import Document

from ..aliyun import convert_office, recognize_image
from ..storage import delete, list_keys, read_bytes

logger = logging.getLogger("oa-worker.parser")

ALLOWED_MAGIC = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/zip",
}


@dataclass
class ParsedPage:
    page: int
    text: str
    ocr_used: bool
    confidence: float


def _ocr_image(image_path: Path) -> tuple[str, float]:
    """Recognise a rasterised page through Aliyun OCR."""
    return recognize_image(image_path.read_bytes())


def _parse_pdf(path: Path, temp_dir: Path) -> list[ParsedPage]:
    pdf = fitz.open(path)
    pages: list[ParsedPage] = []
    for index, page in enumerate(pdf):
        text = page.get_text("text").strip()
        density = len(text) / max(1, page.rect.width * page.rect.height / 10_000)
        if len(text) < 80 or density < 1.2:
            image_path = temp_dir / f"page-{index + 1:04d}.png"
            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(image_path)
            ocr_text, confidence = _ocr_image(image_path)
            if len(ocr_text) > len(text):
                text = ocr_text
                pages.append(ParsedPage(index + 1, text, True, confidence))
                continue
        pages.append(ParsedPage(index + 1, text, False, 0.98 if text else 0.0))
    pdf.close()
    return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    document = Document(path)
    text: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text.append("\t".join(cell.text for cell in row.cells))
    for section in document.sections:
        text.extend(p.text for p in section.header.paragraphs if p.text.strip())
        text.extend(p.text for p in section.footer.paragraphs if p.text.strip())
    return [ParsedPage(1, "\n".join(text), False, 0.98)]


def _parse_legacy_doc(storage_key: str) -> list[ParsedPage]:
    """Extract text from a legacy .doc through Aliyun IMM.

    IMM converts the object that is already in OSS, so the file never has to be
    parsed by an in-container office suite.
    """
    prefix = f"{storage_key.rsplit('/', 1)[0]}/imm-doc/{uuid4().hex}/"
    try:
        convert_office(
            storage_key,
            prefix,
            "txt",
            source_type="doc",
            options={"long_text": True, "hold_line_feed": True},
        )
        produced = sorted(key for key in list_keys(prefix) if key.lower().endswith(".txt"))
        if not produced:
            raise RuntimeError("DOC_CONVERSION_FAILED:imm-output-missing")
        text = "\n".join(read_bytes(key).decode("utf-8", errors="replace") for key in produced).strip()
    finally:
        for key in list_keys(prefix):
            try:
                delete(key)
            except Exception:  # cleanup is best effort
                logger.warning("imm_cleanup_failed key=%s", key)
    return [ParsedPage(1, text, False, 0.9 if text else 0.0)]


def parse_document(path: Path, declared_mime: str, storage_key: str) -> dict:
    detected = magic.from_file(str(path), mime=True)
    if detected not in ALLOWED_MAGIC:
        raise RuntimeError("FILE_TYPE_UNSUPPORTED")
    with tempfile.TemporaryDirectory(prefix="parse-") as folder:
        temp_dir = Path(folder)
        if detected == "application/pdf":
            pages = _parse_pdf(path, temp_dir)
        elif declared_mime == "application/msword" or detected == "application/msword":
            pages = _parse_legacy_doc(storage_key)
        else:
            pages = _parse_docx(path)
    full_text = "\n\n".join(page.text for page in pages)
    quality = sum(page.confidence for page in pages) / max(1, len(pages))
    metadata = extract_metadata(full_text)
    return {
        "schemaVersion": "1.0",
        "pageCount": len(pages),
        "isScanned": any(page.ocr_used for page in pages),
        "textQuality": round(quality, 3),
        "pages": [{"page": p.page, "text": p.text, "ocrUsed": p.ocr_used, "confidence": round(p.confidence, 3)} for p in pages],
        "sections": [], "tables": [], "metadataCandidates": metadata, "fullText": full_text,
    }


def extract_metadata(text: str) -> dict:
    application = re.search(r"(?:申请号|Application\s*No\.?)[\s：:]*([A-Z]{0,2}\d[\d\.]+)", text, re.I)
    serial = re.search(r"(?:通知书编号|发文序号)[\s：:]*([^\s\n]+)", text)
    issue_date = re.search(r"(?:发文日|日期)[\s：:]*(\d{4}[-年]\d{1,2}[-月]\d{1,2}日?)", text)
    cited = []
    # OCR emits each recognised line separately, so the label and its publication
    # number often end up on different lines; allow the gap to contain newlines
    # while still bounding it to 80 characters.
    for label, publication in re.findall(r"\b(D\d+)\b[\s\S]{0,80}?\b((?:CN|EP|US|WO)\s?\d+[A-Z]\d?)\b", text, re.I):
        cited.append({"label": label.upper(), "publicationNo": publication.replace(" ", "").upper()})
    return {
        "applicationNo": application.group(1) if application else None,
        "serialNo": serial.group(1) if serial else None,
        "issueDate": issue_date.group(1) if issue_date else None,
        "citedPriorArts": cited,
    }
