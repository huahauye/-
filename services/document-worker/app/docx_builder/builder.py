from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FORBIDDEN = ["OPENAI_API_KEY", "AI认为", "Chain-of-Thought", "内部审核", "定稿前核验事项"]


def _font(run, name: str = "宋体", size: float = 12, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def _page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    _font(run, size=9)


def _paragraph(document, value: dict):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if value.get("kind") == "body" else None
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(value["text"])
    _font(run, name="仿宋" if value.get("kind") == "quote" else "宋体", bold=value.get("kind") == "subheading")


def build_statement(statement: dict, destination: Path) -> None:
    if statement.get("reviewNotes"):
        raise RuntimeError("DOCX_BUILD_FAILED:reviewNotes must be empty")
    serialized = str(statement)
    if any(term.lower() in serialized.lower() for term in FORBIDDEN):
        raise RuntimeError("DOCX_BUILD_FAILED:internal text detected")
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54); section.bottom_margin = Cm(2.54); section.left_margin = Cm(2.8); section.right_margin = Cm(2.6)
    _page_number(section.footer.paragraphs[0])
    title = document.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run("意 见 陈 述 书"), name="黑体", size=18, bold=True)
    meta = document.add_table(rows=0, cols=2); meta.autofit = False
    for label, value in [("申请号", statement.get("applicationNumber", "")), ("发明名称", statement.get("inventionTitle", "")), ("申请人", statement.get("applicant", "")), ("审查意见", f"{statement.get('oaRound', '')}  {statement.get('oaDate', '')}")]:
        cells = meta.add_row().cells
        _font(cells[0].paragraphs[0].add_run(label), bold=True)
        _font(cells[1].paragraphs[0].add_run(str(value)))
    document.add_paragraph()
    for paragraph in statement.get("lead", []): _paragraph(document, paragraph)
    for item in statement.get("sections", []):
        heading = document.add_paragraph(); heading.paragraph_format.keep_with_next = True
        _font(heading.add_run(item["heading"]), name="黑体", size=13, bold=True)
        for paragraph in item.get("paragraphs", []): _paragraph(document, paragraph)
    for paragraph in statement.get("closing", []): _paragraph(document, paragraph)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    check = Document(destination)
    text_parts = [paragraph.text for paragraph in check.paragraphs]
    text_parts.extend(
        paragraph.text
        for table in check.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    body = "\n".join(text_parts)
    application_number = statement.get("applicationNumber", "")
    if not body.strip() or (application_number and application_number not in body):
        raise RuntimeError("DOCX_BUILD_FAILED:post-build validation")
