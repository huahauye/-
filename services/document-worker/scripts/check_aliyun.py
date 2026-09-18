#!/usr/bin/env python
"""Smoke-test the Aliyun wiring from inside the worker image.

Run from ``services/document-worker``:

    python scripts/check_aliyun.py                  # OSS round-trip only
    python scripts/check_aliyun.py --imm            # + IMM DOCX -> PDF conversion
    python scripts/check_aliyun.py --ocr page.png   # + OCR on a local image
    python scripts/check_aliyun.py --ocr page.png --raw-ocr

Every check prints ``PASS`` or ``FAIL`` plus the underlying error code, so a
misconfiguration is visible before a real case is uploaded.
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import settings  # noqa: E402
from app.storage import client, delete, download, list_keys, read_bytes, upload, upload_bytes  # noqa: E402


def _mask(value: str) -> str:
    if not value:
        return "<unset>"
    return f"{value[:6]}…({len(value)} chars)"


def show_config() -> None:
    cfg = settings()
    print("resolved configuration")
    print(f"  s3_endpoint        = {cfg.s3_endpoint}")
    print(f"  s3_region          = {cfg.s3_region}")
    print(f"  s3_bucket          = {cfg.s3_bucket}")
    print(f"  s3_signature       = {cfg.s3_signature_version}  force_path_style={cfg.s3_force_path_style}")
    print(f"  s3_access_key_id   = {_mask(cfg.s3_access_key_id)}")
    print(f"  imm_project        = {cfg.aliyun_imm_project or '<unset>'}")
    print(f"  imm_endpoint       = {cfg.imm_endpoint_resolved}")
    print(f"  imm_access_key_id  = {_mask(cfg.imm_key_id)}")
    print(f"  ocr_endpoint       = {cfg.ocr_endpoint_resolved}  mode={cfg.aliyun_ocr_mode}")
    print(f"  ocr_access_key_id  = {_mask(cfg.ocr_key_id)}")
    print()


def check_oss() -> bool:
    cfg = settings()
    key = f"diagnostics/oss-check/{uuid4().hex}.txt"
    payload = b"oa-agent oss check"
    try:
        upload_bytes(payload, key, "text/plain; charset=utf-8")
        client().head_object(Bucket=cfg.s3_bucket, Key=key)
        assert read_bytes(key) == payload, "downloaded bytes differ from uploaded bytes"
        delete(key)
    except Exception as error:  # noqa: BLE001 - the whole point is to report anything
        print(f"FAIL  oss round-trip: {type(error).__name__}: {error}")
        return False
    print("PASS  oss round-trip (put / head / get / delete)")
    return True


def check_imm() -> bool:
    from docx import Document

    from app.aliyun import convert_office, oss_uri

    cfg = settings()
    token = uuid4().hex
    source_key = f"diagnostics/imm-check/{token}/sample.docx"
    prefix = f"diagnostics/imm-check/{token}/output/"
    local_root = Path(cfg.worker_temp_dir) / f"imm-check-{token}"
    local_root.mkdir(parents=True, exist_ok=True)
    docx_path = local_root / "sample.docx"
    pdf_path = local_root / "sample.pdf"

    try:
        document = Document()
        document.add_heading("意见陈述书转换自检", level=1)
        document.add_paragraph("本文件用于验证阿里云 IMM 的 Office 转换链路。")
        document.save(docx_path)

        upload(docx_path, source_key, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        print(f"      created task for {oss_uri(source_key)}")
        task_id = convert_office(source_key, prefix, "pdf")
        print(f"      imm task {task_id} succeeded")

        produced = [key for key in list_keys(prefix) if key.lower().endswith(".pdf")]
        if not produced:
            raise RuntimeError("IMM produced no PDF under the target prefix")
        download(produced[0], pdf_path)
        head = pdf_path.read_bytes()[:5]
        if head != b"%PDF-":
            raise RuntimeError(f"downloaded object is not a PDF (starts with {head!r})")
        print(f"PASS  imm office conversion ({pdf_path.stat().st_size} bytes, {produced[0]})")
        return True
    except Exception as error:  # noqa: BLE001
        print(f"FAIL  imm office conversion: {type(error).__name__}: {error}")
        print("      Reminder: the IMM project must live in the same region as the bucket,")
        print("      the RAM user needs imm:CreateOfficeConversionTask and imm:GetTask,")
        print("      and OSS access to the bucket. Also enable the IMM service first.")
        return False
    finally:
        for key in list_keys(prefix) + [source_key]:
            try:
                delete(key)
            except Exception:  # noqa: BLE001
                pass


def check_ocr(image_path: Path, dump_raw: bool) -> bool:
    from app.aliyun import recognize_image

    try:
        data = image_path.read_bytes()
    except OSError as error:
        print(f"FAIL  ocr: cannot read {image_path}: {error}")
        return False

    if dump_raw:
        try:
            from alibabacloud_ocr_api20210707 import models as ocr_models

            from app.aliyun import ocr_client

            response = ocr_client().recognize_advanced(ocr_models.RecognizeAdvancedRequest(body=io.BytesIO(data)))
            print("      raw OCR payload:")
            print(json.dumps(json.loads(response.body.data or "{}"), ensure_ascii=False, indent=2)[:4000])
        except Exception as error:  # noqa: BLE001
            print(f"FAIL  ocr raw payload: {type(error).__name__}: {error}")
            return False

    try:
        text, confidence = recognize_image(data)
    except Exception as error:  # noqa: BLE001
        print(f"FAIL  ocr: {type(error).__name__}: {error}")
        return False
    if not text:
        print("FAIL  ocr: recognised no text (check 文字识别 OCR is activated and the key is authorised)")
        return False
    print(f"PASS  ocr ({len(text)} chars, confidence={confidence:.3f})")
    print("      " + text[:200].replace("\n", " / "))
    return True


def main() -> int:
    parser = argparse.ArgumentParser(description="Aliyun wiring smoke test")
    parser.add_argument("--imm", action="store_true", help="also run a DOCX -> PDF conversion through IMM")
    parser.add_argument("--ocr", type=Path, help="also run OCR on this local PNG/JPEG file")
    parser.add_argument("--raw-ocr", action="store_true", help="print the raw OCR response payload")
    args = parser.parse_args()

    show_config()
    results = [check_oss()]
    if args.imm:
        results.append(check_imm())
    if args.ocr:
        results.append(check_ocr(args.ocr, args.raw_ocr))

    print()
    if all(results):
        print("all checks passed")
        return 0
    print("some checks failed")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
